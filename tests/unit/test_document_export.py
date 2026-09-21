"""
Generated-document export (2026-09-20): markdown → docx/pdf/odt/rtf/html/txt.

Drives THE deployed converters (real pandoc / LibreOffice / fpdf2 when present —
a converter that is absent skips its own case, the ladder tests still run).
"""

import shutil
import zipfile
from types import SimpleNamespace

import pytest

import knowledge.document_export as dx

MD = (
    '---\ntitle: "Fair Resume"\ntype: draft\n---\n'
    "# Luke Example\n\n**Actuarial Analyst** — Mercer, 2022\n\n"
    "- Automated 100+ reports\n- Maintained rate models\n\n## Skills\n\nPython, SQL\n"
)


@pytest.fixture
def md_file(tmp_path):
    p = tmp_path / "fair-resume.md"
    p.write_text(MD, encoding="utf-8")
    return p


@pytest.mark.parametrize("text,expected", [
    ("write a new document as a .docx", "docx"),
    ("please save it as docx", "docx"),
    ("give me a PDF of that", "pdf"),
    ("make it a Word document", "docx"),
    ("export to plain text", "txt"),
    ("use fair_resume.docx and write a pdf", "pdf"),      # filename names the INPUT
    ("write a new document using info in fair_resume.docx", ""),
    ("not a pdf, a docx please", "docx"),                 # negation-aware
    ("write a report about climate change", ""),
    ("a word about formatting", ""),
    ("", ""),
])
def test_detect_requested_format(text, expected):
    assert dx.detect_requested_format(text) == expected


def test_frontmatter_never_reaches_the_export(md_file):
    out = dx.export_document(md_file, "txt")
    text = out.read_text(encoding="utf-8")
    assert "type: draft" not in text and "Luke Example" in text and "**" not in text


def test_html_export(md_file):
    html = dx.export_document(md_file, "html").read_text(encoding="utf-8")
    assert "<h1>Luke Example</h1>" in html and "<li>Automated 100+ reports</li>" in html


@pytest.mark.skipif(not shutil.which("pandoc"), reason="pandoc not installed")
@pytest.mark.parametrize("fmt", ["docx", "odt", "rtf"])
def test_pandoc_formats_write_real_files(md_file, fmt):
    out = dx.export_document(md_file, fmt)
    assert out.suffix == f".{fmt}" and out.stat().st_size > 200
    if fmt == "docx":
        with zipfile.ZipFile(out) as z:
            body = z.read("word/document.xml").decode("utf-8")
        assert "Luke Example" in body and "Automated 100+ reports" in body
        assert "type: draft" not in body


def test_pdf_via_fpdf_fallback_is_a_pdf_with_unicode(md_file, monkeypatch):
    pytest.importorskip("fpdf")
    monkeypatch.setattr(dx, "_soffice_pdf",
                        lambda *a, **k: (_ for _ in ()).throw(dx.DocumentExportError("no soffice")))
    out = dx.export_document(md_file, "pdf")  # body carries an em-dash
    assert out.read_bytes()[:5] == b"%PDF-" and out.stat().st_size > 500


@pytest.mark.slow
@pytest.mark.skipif(not (shutil.which("pandoc") and shutil.which("soffice")),
                    reason="pandoc + LibreOffice not installed")
def test_pdf_via_libreoffice(md_file):
    out = dx.export_document(md_file, "pdf")
    assert out.read_bytes()[:5] == b"%PDF-"


def test_never_overwrites(md_file):
    first = dx.export_document(md_file, "txt")
    second = dx.export_document(md_file, "txt")
    assert first != second and second.name == "fair-resume-2.txt" and first.exists()


def test_ladder_falls_through_then_fails_honestly(md_file, monkeypatch):
    def boom(*a, **k):
        raise dx.DocumentExportError("pandoc is not installed")
    monkeypatch.setattr(dx, "_pandoc", boom)
    monkeypatch.setattr(dx, "_python_docx", boom)
    with pytest.raises(dx.DocumentExportError, match="could not write .docx"):
        dx.export_document(md_file, "docx")
    assert not md_file.with_suffix(".docx").exists()
    with pytest.raises(dx.DocumentExportError, match="unsupported"):
        dx.export_document(md_file, "exe")


# --- handler wiring, through THE deployed _run_doc_generation -----------------

def _handler_ctx(user_text, md_path):
    import gui.handlers as handlers
    import knowledge.document_generator as dg_mod

    class FakeGenerator:
        def __init__(self, **kwargs):
            pass

        def repoint_index(self, old, new):
            FakeGenerator.repointed = (str(old), str(new))

        async def generate(self, **kwargs):
            return SimpleNamespace(title="T", path=str(md_path), doc_type="report",
                                   sources=[], sections_count=1, word_count=10)

    ctx = SimpleNamespace(
        orchestrator=SimpleNamespace(
            prompt_builder=None, memory_system=None,
            model_manager=SimpleNamespace(get_active_model_name=lambda: "m")),
        doc_gen_intent={"topic": "t", "doc_type": "report", "focus": None},
        user_text=user_text, merged_input=user_text, history=[], handled=False,
    )
    return handlers, dg_mod, FakeGenerator, ctx


@pytest.mark.asyncio
async def test_handler_exports_requested_format(md_file, monkeypatch):
    handlers, dg_mod, fake, ctx = _handler_ctx("write a report on rates as html", md_file)
    monkeypatch.setattr(dg_mod, "DocumentGenerator", fake)
    monkeypatch.setattr(handlers, "_write_turn_telemetry", lambda *a, **k: None)
    monkeypatch.setattr(handlers, "_get_session_id", lambda *a: "s1")
    chunks = [c async for c in handlers._run_doc_generation(ctx)]
    assert md_file.with_suffix(".html").exists()
    # One file in the requested format: the intermediate .md is gone, the
    # reply's Path and the index row both follow the export.
    assert not md_file.exists()
    assert "fair-resume.html" in chunks[-1]["content"] and "fair-resume.md" not in chunks[-1]["content"]
    assert fake.repointed == (str(md_file), str(md_file.with_suffix(".html")))


@pytest.mark.asyncio
async def test_handler_reports_export_failure_and_keeps_markdown(md_file, monkeypatch):
    handlers, dg_mod, fake, ctx = _handler_ctx("write a report on rates as a pdf", md_file)
    monkeypatch.setattr(dg_mod, "DocumentGenerator", fake)
    monkeypatch.setattr(handlers, "_write_turn_telemetry", lambda *a, **k: None)
    monkeypatch.setattr(handlers, "_get_session_id", lambda *a: "s1")

    def boom(*a, **k):
        raise dx.DocumentExportError("could not write .pdf: no converter")
    monkeypatch.setattr(dx, "export_document", boom)
    chunks = [c async for c in handlers._run_doc_generation(ctx)]
    assert ctx.handled is True
    assert "not written" in chunks[-1]["content"] and str(md_file) in chunks[-1]["content"]


@pytest.mark.skipif(not shutil.which("pandoc"), reason="pandoc not installed")
def test_stacked_lines_stay_separate_lines_in_docx(tmp_path):
    """Live probe 2026-09-20: a composed resume puts 'School — degree' and
    'Coursework: …' on consecutive lines; plain gfm merged them into one run."""
    p = tmp_path / "r.md"
    p.write_text("Georgia Tech — M.S. Analytics\nCoursework: Simulation\n", encoding="utf-8")
    with zipfile.ZipFile(dx.export_document(p, "docx")) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert "<w:br" in xml
    assert "<br" in dx.export_document(p, "html").read_text(encoding="utf-8")


def test_repoint_index_follows_the_export(tmp_path):
    import json
    import knowledge.document_generator as dg_mod
    gen = dg_mod.DocumentGenerator(model_manager=None, output_dir=str(tmp_path))
    doc = gen.save_prewritten("# Name\n\nBody.\n", topic="fair resume", doc_type="draft")
    gen.repoint_index(doc.path, doc.path.replace(".md", ".docx"))
    rows = json.loads((tmp_path / "index.json").read_text())
    assert len(rows) == 1 and rows[0]["path"].endswith("fair-resume-" + doc.path.rsplit("fair-resume-", 1)[1].replace(".md", ".docx"))


@pytest.mark.skipif(not shutil.which("pandoc"), reason="pandoc not installed")
def test_docx_uses_compact_reference_styles(md_file):
    pytest.importorskip("docx")
    import config.app_config as app_config
    from docx import Document
    doc = Document(str(dx.export_document(md_file, "docx")))
    assert round(doc.sections[0].left_margin.inches, 2) == app_config.DOCUMENT_EXPORT_MARGIN_IN
    normal = doc.styles["Normal"].font
    assert (normal.name == app_config.DOCUMENT_EXPORT_FONT
            and normal.size.pt == app_config.DOCUMENT_EXPORT_BODY_PT)


# --- reference_doc (2026-09-20, template-cast documents): a second attachment
# supplied as a LAYOUT TEMPLATE styles the export directly. ------------------

@pytest.mark.skipif(not shutil.which("pandoc"), reason="pandoc not installed")
def test_user_template_wins_over_the_built_in_reference(md_file, tmp_path):
    pytest.importorskip("docx")
    from docx import Document
    from docx.shared import Inches, Pt

    template_path = tmp_path / "ats_template.docx"
    template = Document()
    template.sections[0].left_margin = Inches(1.25)
    template.sections[0].right_margin = Inches(1.25)
    normal = template.styles["Normal"].font
    normal.name = "Georgia"
    normal.size = Pt(9)
    template.save(str(template_path))

    out = dx.export_document(md_file, "docx", reference_doc=template_path)
    doc = Document(str(out))
    assert round(doc.sections[0].left_margin.inches, 2) == 1.25
    assert doc.styles["Normal"].font.name == "Georgia"
    assert doc.styles["Normal"].font.size.pt == 9


@pytest.mark.skipif(not shutil.which("pandoc"), reason="pandoc not installed")
def test_corrupt_template_falls_back_and_still_succeeds(md_file, tmp_path):
    """A malformed user template must never lose the document — the ladder
    retries once with the built-in reference."""
    bad_template = tmp_path / "bad_template.docx"
    bad_template.write_bytes(b"not a docx")
    out = dx.export_document(md_file, "docx", reference_doc=bad_template)
    assert out.exists() and out.stat().st_size > 0
    with zipfile.ZipFile(out) as z:
        body = z.read("word/document.xml").decode("utf-8")
    assert "Luke Example" in body


def test_reference_doc_missing_file_is_ignored(md_file, tmp_path):
    pytest.importorskip("docx")
    missing = tmp_path / "does_not_exist.docx"
    out = dx.export_document(md_file, "docx", reference_doc=missing)
    assert out.exists() and out.stat().st_size > 0


def test_reference_doc_wrong_format_contributes_nothing(md_file, tmp_path):
    """A template in another format (e.g. .pdf) is structure-text only at the
    compose stage — it must never error out the export."""
    pytest.importorskip("docx")
    other_format = tmp_path / "template.pdf"
    other_format.write_bytes(b"%PDF-1.4 not a real pdf")
    out = dx.export_document(md_file, "docx", reference_doc=other_format)
    assert out.exists() and out.stat().st_size > 0


# --- direct template build (2026-09-20): pandoc's --reference-doc only
# carries named STYLES across, and post-hoc casting still fights the format
# (no numbering ids, wrong heading look, lost spacing/borders). The builder
# below deep-copies the template's OWN exemplar <w:p> elements instead.
# Roles are learned purely structurally (first occurrence, position/format)
# so a memo/cover-letter template builds the same way a resume template does.

_BUILD_MD = (
    "# Name\nHeadline\nContact\n\n"
    "## SECTION\n\n"
    "**Title — Org** | 2022–2024\n*Context line.*\n- bullet **one**\n\n"
    "Plain body.\n"
)

_TEMPLATE_PLACEHOLDER_TEXTS = (
    "Template Placeholder Name", "Template Placeholder Role",
    "TEMPLATE PLACEHOLDER SECTION", "Template Placeholder Entry",
    "Template Placeholder Date", "Template placeholder bullet",
    "Template placeholder body sentence.",
)


def _build_direct_template(path):
    """A template whose look lives on RUNS, not named styles — matching the
    real owner template's shape (Google-Docs export): centered bold-22pt
    title, ONE centered 9.5pt preamble line, a "Heading 2"-styled paragraph
    whose run carries an explicit bold/size/RGB color, a bold+gray entry
    line with a RIGHT tab stop and NO literal tab character in its text (the
    owner's template lost its tab glyph on export), a 10pt List Bullet line,
    and a 10pt body line with space_after set."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Inches, Pt, RGBColor

    tpl = Document()
    tpl.sections[0].left_margin = Inches(0.6)
    tpl.sections[0].right_margin = Inches(0.6)

    title = tpl.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Template Placeholder Name")
    run.bold = True
    run.font.size = Pt(22)

    preamble = tpl.add_paragraph()
    preamble.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = preamble.add_run("Template Placeholder Role")
    run.font.size = Pt(9.5)

    heading = tpl.add_heading("", level=2)
    run = heading.add_run("TEMPLATE PLACEHOLDER SECTION")
    run.bold = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)

    entry = tpl.add_paragraph()
    bold_run = entry.add_run("Template Placeholder Entry")
    bold_run.bold = True
    bold_run.font.size = Pt(10)
    gray_run = entry.add_run("Template Placeholder Date")
    gray_run.font.size = Pt(10)
    gray_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    entry.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)

    bullet = tpl.add_paragraph("Template placeholder bullet", style="List Bullet")
    for r in bullet.runs:
        r.font.size = Pt(10)

    body = tpl.add_paragraph()
    run = body.add_run("Template placeholder body sentence.")
    run.font.size = Pt(10)
    body.paragraph_format.space_after = Pt(6)

    tpl.save(str(path))
    return path


def _numid_of(paragraph):
    """numId of a list paragraph, resolved from its own direct pPr/numPr or
    (python-docx's built-in "List Bullet" style puts numPr on the STYLE, not
    the paragraph) from its style's basedOn chain."""
    from docx.oxml.ns import qn

    def _from_ppr(ppr):
        if ppr is None:
            return None
        num_pr = ppr.find(qn("w:numPr"))
        if num_pr is None:
            return None
        num_id = num_pr.find(qn("w:numId"))
        return num_id.get(qn("w:val")) if num_id is not None else None

    direct = _from_ppr(paragraph._p.pPr)
    if direct is not None:
        return direct
    style = paragraph.style
    seen = set()
    while style is not None and id(style) not in seen:
        seen.add(id(style))
        val = _from_ppr(getattr(style.element, "pPr", None))
        if val is not None:
            return val
        style = getattr(style, "base_style", None)
    return None


@pytest.mark.skipif(not shutil.which("pandoc"), reason="pandoc not installed")
def test_build_from_template_applies_template_look(tmp_path):
    pytest.importorskip("docx")
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import RGBColor

    template_path = _build_direct_template(tmp_path / "template.docx")
    md_path = tmp_path / "resume.md"
    md_path.write_text(_BUILD_MD, encoding="utf-8")

    out = dx.export_document(md_path, "docx", reference_doc=template_path)
    doc = Document(str(out))
    tpl_doc = Document(str(template_path))

    full_text = "\n".join(p.text for p in doc.paragraphs)
    for placeholder in _TEMPLATE_PLACEHOLDER_TEXTS:
        assert placeholder not in full_text

    assert len(doc.paragraphs) == 8

    title_para = doc.paragraphs[0]
    assert title_para.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title_para.runs[0].font.size.pt == 22
    assert title_para.runs[0].bold is True

    # ONE preamble exemplar in the template, TWO preamble lines in the
    # markdown — the second reuses the last (only) learned exemplar.
    preamble1, preamble2 = doc.paragraphs[1], doc.paragraphs[2]
    assert preamble1.text == "Headline" and preamble2.text == "Contact"
    for p in (preamble1, preamble2):
        assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert p.runs[0].font.size.pt == 9.5

    heading_para = doc.paragraphs[3]
    assert heading_para.text == "SECTION"
    heading_run = heading_para.runs[0]
    assert heading_run.bold is True
    assert heading_run.font.size.pt == 9.5
    assert heading_run.font.color.rgb == RGBColor(0x2B, 0x2B, 0x2B)

    entry_para = doc.paragraphs[4]
    assert "\t" in entry_para.text
    tab_stops = list(entry_para.paragraph_format.tab_stops)
    assert any(ts.alignment == WD_TAB_ALIGNMENT.RIGHT for ts in tab_stops)
    bold_entry_run = next(r for r in entry_para.runs if r.bold)
    assert bold_entry_run.text.strip() == "Title — Org"
    date_run = next(r for r in entry_para.runs if "2022" in r.text)
    assert not date_run.bold

    context_para = doc.paragraphs[5]
    assert context_para.text == "Context line."
    assert context_para.runs[0].italic is True
    assert context_para.runs[0].font.size.pt == 10

    bullet_para = doc.paragraphs[6]
    assert bullet_para.text == "bullet one"
    tpl_bullet = next(p for p in tpl_doc.paragraphs if p.text == "Template placeholder bullet")
    assert _numid_of(bullet_para) is not None
    assert _numid_of(bullet_para) == _numid_of(tpl_bullet)
    bold_bullet_run = next(r for r in bullet_para.runs if r.bold)
    assert bold_bullet_run.text == "one"

    body_para = doc.paragraphs[7]
    assert body_para.text == "Plain body."
    tpl_body = next(p for p in tpl_doc.paragraphs if p.text == "Template placeholder body sentence.")
    assert body_para.paragraph_format.space_after == tpl_body.paragraph_format.space_after

    assert doc.sections[0].left_margin == tpl_doc.sections[0].left_margin
    assert doc.sections[0].right_margin == tpl_doc.sections[0].right_margin


@pytest.mark.skipif(not shutil.which("pandoc"), reason="pandoc not installed")
def test_build_from_template_memo_case(tmp_path):
    """Structural, not resume-specific: a memo builds the same way."""
    pytest.importorskip("docx")
    template_path = _build_direct_template(tmp_path / "template.docx")
    md_path = tmp_path / "memo.md"
    md_path.write_text("# Memo\nTo: X\n\n## Background\n\nText.\n", encoding="utf-8")

    out = dx.export_document(md_path, "docx", reference_doc=template_path)
    assert out.exists() and out.stat().st_size > 0


@pytest.mark.skipif(not shutil.which("pandoc"), reason="pandoc not installed")
def test_build_from_template_failure_falls_back_to_pandoc_cast(md_file, tmp_path, monkeypatch):
    """A raise from the direct builder (even a fully-replaced one) must
    never lose the document — the ladder falls back to pandoc(+cast)."""
    pytest.importorskip("docx")
    template_path = _build_direct_template(tmp_path / "template.docx")

    def boom(*a, **k):
        raise RuntimeError("boom")

    monkeypatch.setattr(dx, "_build_from_template", boom)
    out = dx.export_document(md_file, "docx", reference_doc=template_path)
    assert out.exists() and out.stat().st_size > 0
    with zipfile.ZipFile(out) as z:
        body = z.read("word/document.xml").decode("utf-8")
    assert "Luke Example" in body


@pytest.mark.skipif(not shutil.which("pandoc"), reason="pandoc not installed")
def test_cast_template_styles_failure_still_yields_valid_docx(md_file, tmp_path, monkeypatch):
    """The cast is best-effort — a raise (even from a fully-replaced
    _cast_template_styles) must never lose the document pandoc wrote."""
    pytest.importorskip("docx")
    from docx import Document

    template_path = tmp_path / "template.docx"
    Document().save(str(template_path))

    def boom(*a, **k):
        raise RuntimeError("boom")

    monkeypatch.setattr(dx, "_cast_template_styles", boom)
    out = dx.export_document(md_file, "docx", reference_doc=template_path)
    assert out.exists() and out.stat().st_size > 0
    Document(str(out))  # must open without error


def test_cast_template_styles_no_python_docx_is_a_noop(md_file, tmp_path, monkeypatch):
    """python-docx missing → warn and leave the pandoc output untouched."""
    import builtins

    real_import = builtins.__import__

    def fake_import(name, *a, **k):
        if name == "docx" or name.startswith("docx."):
            raise ImportError("no docx")
        return real_import(name, *a, **k)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    out_path = tmp_path / "out.docx"
    out_path.write_bytes(b"placeholder")
    dx._cast_template_styles(out_path, tmp_path / "template.docx")
    assert out_path.read_bytes() == b"placeholder"
