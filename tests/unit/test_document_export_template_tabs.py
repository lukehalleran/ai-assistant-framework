"""Template builder right-tab polish (oct26b D1 item 3; incident 2026-09-27).
class: BC-45

(a) the right-tab split applied to every block containing " | " (headings,
bullets, prose); it is for bold-lead ENTRY lines with a short date-like right
part. (c) a long entry's left part pushed the right part past the margin; it is
shrunk to fit or the line stays inline. Real python-docx template, invented lines.
"""

import pytest

pytest.importorskip("docx")

import knowledge.document_export as dx  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.text import WD_TAB_ALIGNMENT  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402


def _template(path):
    tpl = Document()
    title = tpl.add_paragraph()
    r = title.add_run("Placeholder Name")
    r.bold = True
    r.font.size = Pt(20)
    h = tpl.add_heading("", level=2)
    r = h.add_run("PLACEHOLDER SECTION")
    r.bold = True
    entry = tpl.add_paragraph()
    r = entry.add_run("Placeholder Entry")
    r.bold = True
    r.font.size = Pt(10)
    r = entry.add_run("Placeholder Date")
    r.font.size = Pt(10)
    entry.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    body = tpl.add_paragraph()
    body.add_run("Placeholder body.").font.size = Pt(10)
    tpl.save(str(path))
    return path


def _build(tmp_path, md):
    tpl = _template(tmp_path / "t.docx")
    out = tmp_path / "o.docx"
    dx._build_from_template(md, tpl, out)
    return Document(str(out))


def _has_tab(p):
    return p._p.find(".//" + qn("w:r") + "/" + qn("w:tab")) is not None


def _para(doc, startswith):
    return next(p for p in doc.paragraphs if p.text.lstrip("\t").startswith(startswith) or startswith in p.text)


def test_entry_line_with_date_still_splits(tmp_path):
    doc = _build(tmp_path, "# Rowan Vale\n\n## EXPERIENCE\n\n**Acme Mutual** | Tamsin City | Jan 2020 - Present\n")
    p = _para(doc, "Acme Mutual")
    assert _has_tab(p) and p.text.endswith("Jan 2020 - Present")


@pytest.mark.parametrize("line", [
    "## Skills | Tools",
    "- Led the team | and shipped on time",
    "Plain prose line with a | pipe in the middle 2020",
    "**Acme Mutual** | Led the platform team and shipped",
])
def test_non_entry_or_non_date_lines_stay_inline(tmp_path, line):
    doc = _build(tmp_path, f"# Rowan Vale\n\n## EXPERIENCE\n\n{line}\n")
    p = _para(doc, line.lstrip("#-* ").split(" | ")[0].replace("*", "").strip()[:8])
    assert not _has_tab(p), p.text


def test_long_entry_left_part_is_shrunk_to_fit(tmp_path):
    left = "**Senior Associate Pricing Analyst, Acme Mutual Reinsurance** | Tamsin City, Rowan Province"
    doc = _build(tmp_path, f"# Rowan Vale\n\n## EXPERIENCE\n\n{left} | 2019 - 2023\n")
    p = _para(doc, "Senior Associate")
    assert _has_tab(p)
    sizes = [r.font.size.pt for r in p.runs if r.font.size]
    assert sizes and min(sizes) < 10 and min(sizes) >= 10 * 0.8 - 0.5


def test_unfittable_entry_stays_inline_not_overrun(tmp_path):
    left = "**" + "Extraordinarily Long Employer Name " * 8 + "**"
    doc = _build(tmp_path, f"# Rowan Vale\n\n## EXPERIENCE\n\n{left} | 2019 - 2023\n")
    p = _para(doc, "Extraordinarily")
    assert not _has_tab(p) and p.text.endswith("| 2019 - 2023")
