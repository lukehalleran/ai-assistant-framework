"""
Generated-document export (2026-09-20): markdown → docx / pdf / odt / rtf / html / txt.

DocumentGenerator only ever wrote markdown; "write me a new resume as a .docx"
had no way to produce the file the user can actually open in Word. This module
converts the ALREADY-WRITTEN markdown document (the .md stays on disk as the
source of record) into a sibling file with the requested extension.

Module contract:
  - detect_requested_format(text) -> "" | one of EXPORT_FORMATS
      Closed vocabulary of format names/extensions, negation-aware
      (utils.trigger_match.is_negated), filename-blind ("fair_resume.docx" in
      the message names the INPUT, it is not a request for docx output).
  - export_document(md_path, fmt, *, reference_doc=None) -> Path   (blocking — call via asyncio.to_thread)
      Converter ladder per format; never overwrites (versioned -N suffix);
      raises DocumentExportError with the reason when no converter can run —
      the caller reports it honestly, the .md is still there. reference_doc
      (2026-09-20, template-cast documents) lets a caller-supplied styling
      file (e.g. a second attachment used as a layout template) win over the
      built-in compact reference for docx/odt/pdf.

Ladders:
  docx/odt/rtf : pandoc → (docx only) python-docx minimal renderer
  pdf          : pandoc→docx + soffice --convert-to pdf → fpdf2 write_html
  html         : mistune (pure python)
  txt          : markdown with frontmatter + emphasis/heading markers stripped
"""

from __future__ import annotations

import copy
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from config import app_config
from utils.logging_utils import get_logger
from utils.trigger_match import is_negated

logger = get_logger("document_export")

EXPORT_FORMATS = ("docx", "pdf", "odt", "rtf", "html", "txt")
_CONVERT_TIMEOUT_S = 90

# Extension/format tokens. The lookbehind refuses a token glued to a filename
# ("resume.docx", "notes/out.pdf"); a free-standing ".docx" or "docx" is a request.
_FORMAT_TOKEN_RE = re.compile(r"(?<![\w/.])\.?(docx|pdf|odt|rtf|html|txt)\b", re.I)
# Product names for the same closed set (a format has a handful of names, not an
# open phrase space).
_FORMAT_NAME_RE = re.compile(
    r"\b(?:(?P<docx>(?:ms|microsoft)?\s*word\s+(?:doc(?:ument)?|file|format))"
    r"|(?P<odt>(?:libre|open)\s*office\s+(?:doc(?:ument)?|file|format))"
    r"|(?P<txt>plain[\s-]*text(?:\s+file)?)"
    r"|(?P<html>web\s*page))\b",
    re.I,
)
# The shared matcher scopes VERB negation ("don't make a pdf"); a format is a
# noun, so "not a pdf" / "no pdf" negate it by direct adjacency.
_ADJACENT_NO_RE = re.compile(r"\b(?:not|no)\s+(?:as\s+)?(?:an?\s+)?\.?$", re.I)
_FONT_DIRS = (
    "/usr/share/fonts/dejavu-sans-fonts", "/usr/share/fonts/truetype/dejavu",
    "/usr/share/fonts/TTF", "/Library/Fonts", "C:/Windows/Fonts",
)


class DocumentExportError(RuntimeError):
    """No converter could produce the requested format."""


def detect_requested_format(text: str) -> str:
    """The output format the user asked for, or '' (markdown default)."""
    text = text or ""
    hits: list[tuple[int, str]] = []
    for m in _FORMAT_TOKEN_RE.finditer(text):
        hits.append((m.start(), m.group(1).lower()))
    for m in _FORMAT_NAME_RE.finditer(text):
        hits.append((m.start(), m.lastgroup or ""))
    for start, fmt in sorted(hits):
        if fmt and not is_negated(text, start) and not _ADJACENT_NO_RE.search(text[:start]):
            return fmt
    return ""


def strip_frontmatter(markdown: str) -> str:
    """Drop the generator's leading YAML block — it is index metadata, not content."""
    if markdown.startswith("---"):
        end = markdown.find("\n---", 3)
        if end != -1:
            return markdown[end + 4:].lstrip("\n")
    return markdown


def _versioned(path: Path) -> Path:
    if not path.exists():
        return path
    n = 2
    while True:
        candidate = path.with_name(f"{path.stem}-{n}{path.suffix}")
        if not candidate.exists():
            return candidate
        n += 1


def _run(cmd: list[str]) -> None:
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=_CONVERT_TIMEOUT_S)
    if proc.returncode != 0:
        raise DocumentExportError(f"{cmd[0]} failed: {(proc.stderr or proc.stdout).strip()[:300]}")


def _reference_docx(tmp: str) -> "Path | None":
    """Compact reference styles for pandoc's docx writer, or None.

    pandoc's stock template is 12pt with 1in margins and wide heading gaps — a
    one-page resume exported at ~1.4 pages (live 2026-09-20). Calibri 11 (Carlito
    is its metric twin where Calibri is absent), 0.6in margins, tight spacing,
    black headings — configurable via config.yaml `document_generation.export_*`
    (app_config.DOCUMENT_EXPORT_FONT/BODY_PT/MARGIN_IN, read live so a config
    reload takes effect without restarting the export path). Needs python-docx;
    without it the stock template is used.
    """
    try:
        from docx import Document  # lazy import: startup-cost
        from docx.shared import Inches, Pt, RGBColor  # lazy import: startup-cost
    except ImportError:
        return None
    exe = shutil.which("pandoc")
    ref = Path(tmp) / "reference.docx"
    with open(ref, "wb") as fh:
        proc = subprocess.run([exe, "--print-default-data-file", "reference.docx"],
                              stdout=fh, stderr=subprocess.PIPE, timeout=_CONVERT_TIMEOUT_S)
    if proc.returncode != 0:
        return None
    doc = Document(str(ref))
    margin_in = app_config.DOCUMENT_EXPORT_MARGIN_IN
    font_name = app_config.DOCUMENT_EXPORT_FONT
    body_pt = app_config.DOCUMENT_EXPORT_BODY_PT
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(margin_in)
        section.top_margin = section.bottom_margin = Inches(margin_in)
    sizes = {"Title": 18, "Heading 1": 18, "Heading 2": 12, "Heading 3": 11}
    for style in doc.styles:
        font = getattr(style, "font", None)
        if font is None or style.type != 1:  # paragraph styles only
            continue
        font.name = font_name
        font.size = Pt(sizes.get(style.name, body_pt))
        if style.name in sizes:
            font.color.rgb = RGBColor(0, 0, 0)
            font.bold = True
        fmt = style.paragraph_format
        fmt.space_before = Pt(8 if style.name == "Heading 2" else 0)
        fmt.space_after = Pt(2)
        fmt.line_spacing = 1.0
    doc.save(str(ref))
    return ref


def _para_style_name(paragraph) -> str:
    style = getattr(paragraph, "style", None)
    return (getattr(style, "name", "") or "") if style is not None else ""


def _is_heading_style(paragraph) -> bool:
    return _para_style_name(paragraph).strip().lower().startswith("heading")


def _style_chain_has_numpr(style, qn) -> bool:
    """Walk a style and its basedOn chain for an inherited ``<w:numPr>``.

    python-docx's built-in "List Bullet" style carries numPr on the STYLE
    definition, not on each paragraph's own pPr (verified live: pandoc's own
    list output puts numPr directly on the paragraph; a python-docx-authored
    template paragraph using a named list style does not)."""
    element = getattr(style, "element", None)
    seen: set[int] = set()
    while element is not None and id(element) not in seen:
        seen.add(id(element))
        pPr = getattr(element, "pPr", None)
        if pPr is not None and pPr.find(qn("w:numPr")) is not None:
            return True
        element = getattr(element, "base_style", None)
    return False


def _has_numpr(paragraph, qn) -> bool:
    """True if this paragraph is a list item, checking both the paragraph's
    own direct pPr/numPr (how pandoc's output represents list items) and an
    inherited numPr from its named style (how a python-docx-authored "List
    Bullet" paragraph represents one)."""
    pPr = getattr(paragraph._p, "pPr", None)
    if pPr is not None and pPr.find(qn("w:numPr")) is not None:
        return True
    return _style_chain_has_numpr(paragraph.style, qn)


def _first_nonempty_run(paragraph):
    for run in paragraph.runs:
        if run.text and run.text.strip():
            return run
    return None


def _record_exemplar(paragraph) -> dict:
    """Snapshot one template paragraph's look: paragraph-level alignment/
    spacing/tab-stops, and from its first non-empty run: font name, size,
    bold, italic, color."""
    fmt = paragraph.paragraph_format
    tab_stops = [(ts.position, ts.alignment) for ts in fmt.tab_stops]
    run = _first_nonempty_run(paragraph)
    color = None
    if run is not None:
        try:
            color = run.font.color.rgb
        except Exception:  # degrades: color is cosmetic-only, never applied on output anyway
            color = None
    return {
        "style_name": _para_style_name(paragraph),
        "alignment": paragraph.alignment,
        "space_before": fmt.space_before,
        "space_after": fmt.space_after,
        "line_spacing": fmt.line_spacing,
        "tab_stops": tab_stops,
        "text": paragraph.text,
        "font_name": run.font.name if run is not None else None,
        "font_size": run.font.size if run is not None else None,
        "bold": run.font.bold if run is not None else None,
        "italic": run.font.italic if run is not None else None,
        "color": color,
    }


def _learn_template_roles(paragraphs, qn) -> dict:
    """First-occurrence structural-role PARAGRAPHS, learned from the template.

    Roles: title (first non-empty paragraph), preamble[i] (non-empty
    paragraphs after the title and before the first heading-styled
    paragraph), heading (first heading-styled paragraph), entry (first
    non-list, non-heading paragraph after the first heading whose first
    non-empty run is bold), bullet (first list paragraph anywhere), body
    (first non-list, non-heading, non-bold-lead paragraph after the first
    heading). Purely structural/positional — no document-type keywords, so a
    cover letter or memo template is learned the same way a resume is.

    Shared by `_learn_template_exemplars` (style-dict snapshot, used by the
    pandoc-cast fallback path) and `_build_from_template` (keeps the actual
    python-docx Paragraph so its `<w:p>` element can be deep-copied).
    """
    roles: dict = {"preamble": []}
    title_idx = None
    heading_idx = None
    for idx, paragraph in enumerate(paragraphs):
        if not paragraph.text.strip():
            continue
        if title_idx is None:
            roles["title"] = paragraph
            title_idx = idx
            continue
        if heading_idx is None:
            if _is_heading_style(paragraph):
                roles["heading"] = paragraph
                heading_idx = idx
            else:
                roles["preamble"].append(paragraph)

    for paragraph in paragraphs:
        if _has_numpr(paragraph, qn):
            roles["bullet"] = paragraph
            break

    if heading_idx is not None:
        for paragraph in paragraphs[heading_idx + 1:]:
            if not paragraph.text.strip() or _has_numpr(paragraph, qn) or _is_heading_style(paragraph):
                continue
            run = _first_nonempty_run(paragraph)
            bold_lead = bool(run.font.bold) if run is not None else False
            if bold_lead and "entry" not in roles:
                roles["entry"] = paragraph
            elif not bold_lead and "body" not in roles:
                roles["body"] = paragraph
            if "entry" in roles and "body" in roles:
                break
    return roles


def _learn_template_exemplars(paragraphs, qn) -> dict:
    """Style-dict snapshot per structural role (see `_learn_template_roles`),
    used by the pandoc-cast fallback path (`_cast_template_styles`)."""
    roles = _learn_template_roles(paragraphs, qn)
    exemplars: dict = {"preamble": [_record_exemplar(p) for p in roles["preamble"]]}
    for key in ("title", "heading", "bullet", "entry", "body"):
        if key in roles:
            exemplars[key] = _record_exemplar(roles[key])
    return exemplars


def _classify_output_paragraphs(paragraphs, qn) -> list:
    """Classify every OUTPUT paragraph by the same structural roles: title =
    the Heading 1 paragraph or the first paragraph; preamble = paragraphs
    before the first heading-styled paragraph; entry = bold-lead non-list
    paragraph; bullet = numPr; body = the rest (including italic
    paragraphs — their italic is never touched). Returns a list of
    (paragraph, role, preamble_index) tuples."""
    title_idx = next(
        (i for i, p in enumerate(paragraphs) if _para_style_name(p).strip().lower() == "heading 1"),
        None,
    )
    if title_idx is None:
        title_idx = 0 if paragraphs else None
    heading_idx = None
    if title_idx is not None:
        for i, p in enumerate(paragraphs):
            if i > title_idx and _is_heading_style(p):
                heading_idx = i
                break

    classified = []
    preamble_counter = 0
    for i, p in enumerate(paragraphs):
        if title_idx is not None and i == title_idx:
            classified.append((p, "title", None))
            continue
        if _is_heading_style(p):
            classified.append((p, "heading", None))
            continue
        if not p.text.strip():
            classified.append((p, "skip", None))
            continue
        if heading_idx is not None and i < heading_idx:
            classified.append((p, "preamble", preamble_counter))
            preamble_counter += 1
            continue
        if _has_numpr(p, qn):
            classified.append((p, "bullet", None))
            continue
        run = _first_nonempty_run(p)
        if run is not None and bool(run.font.bold):
            classified.append((p, "entry", None))
        else:
            classified.append((p, "body", None))
    return classified


def _apply_alignment_spacing(paragraph, exemplar: dict) -> None:
    if exemplar.get("alignment") is not None:
        paragraph.alignment = exemplar["alignment"]
    fmt = paragraph.paragraph_format
    if exemplar.get("space_before") is not None:
        fmt.space_before = exemplar["space_before"]
    if exemplar.get("space_after") is not None:
        fmt.space_after = exemplar["space_after"]
    if exemplar.get("line_spacing") is not None:
        fmt.line_spacing = exemplar["line_spacing"]


def _apply_run_font(run, exemplar: dict) -> None:
    """Font name + size only — bold/italic are never touched here; the
    composer's own choices encode the user's instructions."""
    if exemplar.get("font_name"):
        run.font.name = exemplar["font_name"]
    if exemplar.get("font_size") is not None:
        run.font.size = exemplar["font_size"]


def _apply_body_style(paragraph, role: str, exemplar: dict, body_exemplar: "dict | None") -> None:
    """Apply font name/size per run. A paragraph classified "entry" can still
    contain a non-bold run — pandoc's hard_line_breaks stacks a bold
    entry-title line and the very next markdown line (e.g. an italic context
    line) into ONE physical paragraph joined by a line break — such a run
    takes the "body" exemplar instead so it keeps a body look rather than
    inheriting the entry title's, even though it shares the paragraph."""
    for run in paragraph.runs:
        if role == "entry" and body_exemplar is not None and not bool(run.font.bold):
            _apply_run_font(run, body_exemplar)
        else:
            _apply_run_font(run, exemplar)


def _apply_title_cast(paragraph, exemplar: dict) -> None:
    """Keep the paragraph's own style (do not fight over the style name) —
    only cast run-level alignment/size/bold when the exemplar itself is a
    normal-styled paragraph (a Google-Docs-style template puts its title
    LOOK on runs under a generic style, not a named Title/Heading style)."""
    style_name = (exemplar.get("style_name") or "").strip().lower()
    if style_name and not style_name.startswith("normal"):
        return
    _apply_alignment_spacing(paragraph, exemplar)
    for run in paragraph.runs:
        _apply_run_font(run, exemplar)
        if exemplar.get("bold") is not None:
            run.font.bold = exemplar["bold"]


def _apply_entry_tab(paragraph, exemplar: dict, wd_tab_alignment) -> None:
    """Move the LAST " | "-separated segment of an entry line to a right tab
    stop, mirroring a template that puts a tab before its date column — only
    when the exemplar itself has a right-aligned tab stop and its own text
    contains a tab. No date detection: purely "last separator segment goes
    to the right stop"."""
    tab_stops = exemplar.get("tab_stops") or []
    has_right_tab = any(alignment == wd_tab_alignment.RIGHT for _, alignment in tab_stops)
    if not has_right_tab or "\t" not in (exemplar.get("text") or ""):
        return
    sep = " | "
    target_run = None
    for run in reversed(paragraph.runs):
        if sep in (run.text or ""):
            target_run = run
            break
    if target_run is None:
        return
    text = target_run.text
    pos = text.rfind(sep)
    target_run.text = text[:pos] + "\t" + text[pos + len(sep):]
    pf = paragraph.paragraph_format
    for position, alignment in tab_stops:
        pf.tab_stops.add_tab_stop(position, alignment)


def _cast_template_styles(out_docx: Path, template_docx: Path) -> None:
    """Cast the user's own template's structural LOOK onto pandoc's docx output.

    pandoc's ``--reference-doc`` only carries named paragraph STYLES across; a
    template exported from Google Docs (live 2026-09-20: paragraph style
    literally named lowercase "normal") puts its actual look — center
    alignment, run-level bold/size, a right tab stop for a date column — on
    RUNS, which pandoc's style-based mechanism can't see (verified live:
    margins + Heading 2 carried over, the title/entry look did not).

    Structural roles (title/preamble/heading/entry/bullet/body) are learned
    from the template by FIRST OCCURRENCE and matched onto the
    equivalently-classified OUTPUT paragraph — purely positional/formatting
    structure, no document-type keywords, so a cover letter or memo template
    casts the same way a resume does.

    Best-effort only: python-docx is optional, and any failure here must
    never lose the document pandoc already wrote — log a warning and leave
    the file exactly as pandoc produced it.
    """
    try:
        from docx import Document  # lazy import: startup-cost
        from docx.enum.text import WD_TAB_ALIGNMENT  # lazy import: startup-cost
        from docx.oxml.ns import qn  # lazy import: startup-cost
    except ImportError:
        logger.warning("[DocExport] Template style cast skipped: python-docx not installed")
        return

    try:
        template = Document(str(template_docx))
        out = Document(str(out_docx))
        exemplars = _learn_template_exemplars(template.paragraphs, qn)
        preamble_exemplars = exemplars.get("preamble") or []
        body_exemplar = exemplars.get("body")
        classified = _classify_output_paragraphs(out.paragraphs, qn)

        counts: dict[str, int] = {}
        for paragraph, role, preamble_idx in classified:
            if role in ("skip", "heading"):
                continue
            if role == "preamble":
                exemplar = (
                    preamble_exemplars[min(preamble_idx, len(preamble_exemplars) - 1)]
                    if preamble_exemplars else None
                )
            else:
                exemplar = exemplars.get(role)
            if exemplar is None:
                continue
            if role == "title":
                _apply_title_cast(paragraph, exemplar)
            else:
                _apply_alignment_spacing(paragraph, exemplar)
                _apply_body_style(paragraph, role, exemplar, body_exemplar)
                if role == "entry":
                    _apply_entry_tab(paragraph, exemplar, WD_TAB_ALIGNMENT)
            counts[role] = counts.get(role, 0) + 1

        out.save(str(out_docx))
        template_roles = sorted(k for k, v in exemplars.items() if k != "preamble" and v)
        if preamble_exemplars:
            template_roles.append(f"preamble x{len(preamble_exemplars)}")
        logger.info(
            f"[DocExport] Template style cast: template roles={template_roles}; "
            f"output paragraphs restyled={counts}"
        )
    except Exception as e:  # degrades: pandoc's plain reference-doc styling only
        logger.warning(f"[DocExport] Template style cast failed, leaving pandoc output as-is: {e}")


# --- direct template build (2026-09-20) --------------------------------------
# pandoc's --reference-doc only carries named paragraph STYLES across, and the
# post-hoc _cast_template_styles pass above still fights the target format:
# pandoc's own numbering ids don't exist in the user's template (bullets render
# with no glyph), a heading paragraph keeps pandoc's own Heading-2-styled
# fallback look until cast overwrites individual properties, and paragraph
# borders/spacing that live on <w:pPr> never transfer at all. This builds the
# output paragraph-by-paragraph by deep-copying the template's OWN exemplar
# <w:p> elements (bullets/tab stops/run-level colors&sizes/borders/spacing all
# travel for free since the whole element is copied) instead of writing text
# through pandoc first. Best-effort: `export_document` catches ANY failure
# here and falls back to the pandoc(+cast) path — this must never be the only
# way the document can come into being.

_MD_TITLE_RE = re.compile(r"^#\s+(.*)$")
_MD_HEADING_RE = re.compile(r"^#{2,3}\s+(.*)$")
_MD_BULLET_RE = re.compile(r"^[-*+]\s+(.*)$")
_MD_NUM_BULLET_RE = re.compile(r"^\d+\.\s+(.*)$")

# Bold markers (**/__) must be tried before the single-char italic markers so
# "**x**" isn't read as two adjacent italic spans; regex alternation tries
# alternatives left-to-right at each position, so order here is the priority.
_INLINE_SPAN_RE = re.compile(
    r"\*\*(?P<b1>.+?)\*\*"
    r"|__(?P<b2>.+?)__"
    r"|\*(?P<i1>.+?)\*"
    r"|_(?P<i2>.+?)_"
    r"|`(?P<code>[^`]*)`"
    r"|\[(?P<label>[^\]]+)\]\((?P<url>[^)]+)\)"
)


def _parse_inline_spans(text: str) -> list[tuple[str, bool, bool]]:
    """Split one markdown line into (text, bold, italic) segments.

    Backticks are stripped (code gets no distinct run styling here); a
    markdown link renders as its bare label when label == url, else
    "label (url)"."""
    segments: list[tuple[str, bool, bool]] = []
    pos = 0
    for m in _INLINE_SPAN_RE.finditer(text):
        if m.start() > pos:
            segments.append((text[pos:m.start()], False, False))
        if m.group("b1") is not None:
            segments.append((m.group("b1"), True, False))
        elif m.group("b2") is not None:
            segments.append((m.group("b2"), True, False))
        elif m.group("i1") is not None:
            segments.append((m.group("i1"), False, True))
        elif m.group("i2") is not None:
            segments.append((m.group("i2"), False, True))
        elif m.group("code") is not None:
            segments.append((m.group("code"), False, False))
        else:
            label, url = m.group("label"), m.group("url")
            segments.append((label if label == url else f"{label} ({url})", False, False))
        pos = m.end()
    if pos < len(text):
        segments.append((text[pos:], False, False))
    return [s for s in segments if s[0]] or [(text, False, False)]


def _parse_markdown_blocks(markdown_body: str) -> list[dict]:
    """One block per non-empty markdown LINE — hard line breaks become
    separate paragraphs, never merged (unlike the pandoc hard_line_breaks
    path). Roles: title (first `# ` line, or the first line if there is no
    `# ` line anywhere), heading (`##`/`###`), bullet (`-`/`*`/`+`/`1. `),
    else text -> entry (first segment bold) / preamble[i] (before the first
    heading) / body (everything else)."""
    lines = [ln.strip() for ln in markdown_body.splitlines() if ln.strip()]
    has_title_line = any(_MD_TITLE_RE.match(ln) for ln in lines)
    blocks: list[dict] = []
    seen_heading = False
    preamble_idx = 0
    for i, line in enumerate(lines):
        title_m = _MD_TITLE_RE.match(line)
        heading_m = _MD_HEADING_RE.match(line)
        bullet_m = _MD_BULLET_RE.match(line) or _MD_NUM_BULLET_RE.match(line)
        if title_m and has_title_line:
            role, text = "title", title_m.group(1)
        elif not has_title_line and i == 0:
            role, text = "title", line
        elif heading_m:
            role, text = "heading", heading_m.group(1)
            seen_heading = True
        elif bullet_m:
            role, text = "bullet", bullet_m.group(1)
        else:
            role, text = "text", line

        segments = _parse_inline_spans(text)
        if role == "text":
            if segments[0][1]:
                role = "entry"
            elif not seen_heading:
                role = "preamble"
            else:
                role = "body"

        block = {"role": role, "segments": segments}
        if role == "preamble":
            block["preamble_index"] = preamble_idx
            preamble_idx += 1
        blocks.append(block)
    return blocks


def _any_text_exemplar(roles: dict):
    """The most body-like exemplar available, for a role with none of its
    own — the last resort before giving up on a block entirely."""
    if roles.get("body") is not None:
        return roles["body"]
    if roles.get("entry") is not None:
        return roles["entry"]
    pool = roles.get("preamble") or []
    if pool:
        return pool[0]
    for key in ("title", "heading", "bullet"):
        if roles.get(key) is not None:
            return roles[key]
    return None


def _resolve_exemplar(role: str, preamble_index, roles: dict):
    """Fallback ladder: preamble beyond the last learned preamble reuses the
    last one; entry falls back to body; body falls back to entry, then to
    any available text exemplar; heading/title/bullet fall back to body,
    then to any available text exemplar."""
    if role == "preamble":
        pool = roles.get("preamble") or []
        if pool:
            idx = min(preamble_index if preamble_index is not None else 0, len(pool) - 1)
            return pool[idx]
        return _any_text_exemplar(roles)
    if role == "entry":
        return roles.get("entry") or roles.get("body") or _any_text_exemplar(roles)
    if role == "body":
        return roles.get("body") or roles.get("entry") or _any_text_exemplar(roles)
    if role in ("heading", "title", "bullet"):
        return roles.get(role) or roles.get("body") or _any_text_exemplar(roles)
    return roles.get(role) or _any_text_exemplar(roles)


def _has_any_exemplar(roles: dict) -> bool:
    if roles.get("preamble"):
        return True
    return any(roles.get(k) is not None for k in ("title", "heading", "entry", "bullet", "body"))


def _exemplar_run_for(exemplar_paragraph, bold: bool):
    """The exemplar run whose <w:rPr> a new run of this boldness should
    copy: bold segment -> first bold non-empty run, else first non-empty
    run; non-bold segment -> first non-bold non-empty run, else the very
    first run (even if empty)."""
    runs = list(exemplar_paragraph.runs)
    nonempty = [r for r in runs if r.text and r.text.strip()]
    if bold:
        candidates = [r for r in nonempty if bool(r.font.bold)]
        if candidates:
            return candidates[0]
        if nonempty:
            return nonempty[0]
        return runs[0] if runs else None
    candidates = [r for r in nonempty if not bool(r.font.bold)]
    if candidates:
        return candidates[0]
    return runs[0] if runs else None


def _clear_paragraph_runs(p_element, qn) -> None:
    """Delete every run/hyperlink child of a (deep-copied) exemplar <w:p>,
    keeping its <w:pPr> — the paragraph's own look (alignment, spacing,
    borders, tab stops, list numbering) survives; only its text is wiped."""
    for child in list(p_element):
        if child.tag in (qn("w:r"), qn("w:hyperlink")):
            p_element.remove(child)


def _make_run_element(text: str, rpr_source_run, *, bold, italic, qn, leading_tab: bool = False):
    """A fresh <w:r> carrying `text`: <w:rPr> is a deepcopy of
    `rpr_source_run`'s (font/size/color/caps all travel), then bold/italic
    are set explicitly when given (None leaves the copied rPr's own
    bold/italic/caps untouched, for title/heading roles)."""
    from docx.oxml import OxmlElement  # lazy import: startup-cost
    from docx.text.run import Run  # lazy import: startup-cost

    r = OxmlElement("w:r")
    if rpr_source_run is not None:
        src_rpr = rpr_source_run._r.find(qn("w:rPr"))
        if src_rpr is not None:
            r.append(copy.deepcopy(src_rpr))
    if leading_tab:
        r.append(OxmlElement("w:tab"))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    run = Run(r, None)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    return r


def _set_run_text_preserve(run, text: str, qn) -> None:
    run._r.clear_content()
    t = run._r.add_t(text)
    t.set(qn("xml:space"), "preserve")


def _apply_right_tab_split(new_p, exemplar_paragraph, segments, qn, wd_tab_alignment) -> None:
    """If the block's exemplar has a RIGHT tab stop — no literal tab
    character is required in the exemplar's own text, the owner's template
    lost its tab glyph on export — and the rendered line's plain text
    contains " | ", move everything after the LAST " | " into its own run
    preceded by a <w:tab/>, formatted with the exemplar's non-bold run
    properties. No date detection: purely "last separator segment goes to
    the right stop"."""
    from docx.text.paragraph import Paragraph  # lazy import: startup-cost

    tab_stops = exemplar_paragraph.paragraph_format.tab_stops
    if not any(ts.alignment == wd_tab_alignment.RIGHT for ts in tab_stops):
        return
    plain_text = "".join(t for t, _, _ in segments)
    sep = " | "
    if sep not in plain_text:
        return
    paragraph = Paragraph(new_p, None)
    target = None
    for run in reversed(paragraph.runs):
        if sep in (run.text or ""):
            target = run
            break
    if target is None:
        return
    text = target.text
    pos = text.rfind(sep)
    before, after = text[:pos], text[pos + len(sep):]
    nb_source = _exemplar_run_for(exemplar_paragraph, bold=False)
    tab_run = _make_run_element(after, nb_source, bold=None, italic=None, qn=qn, leading_tab=True)
    target._r.addnext(tab_run)
    if before:
        _set_run_text_preserve(target, before, qn)
    else:
        new_p.remove(target._r)


def _append_block_runs(new_p, block: dict, exemplar_paragraph, qn, wd_tab_alignment) -> None:
    role = block["role"]
    keep_exemplar_emphasis = role in ("title", "heading")
    for text, bold, italic in block["segments"]:
        src_run = _exemplar_run_for(exemplar_paragraph, bold)
        r = _make_run_element(
            text, src_run, qn=qn,
            bold=None if keep_exemplar_emphasis else bold,
            italic=None if keep_exemplar_emphasis else italic,
        )
        new_p.append(r)
    _apply_right_tab_split(new_p, exemplar_paragraph, block["segments"], qn, wd_tab_alignment)


def _build_from_template(markdown_body: str, template_docx: Path, out: Path) -> None:
    """Build the docx directly from the template's own exemplar paragraphs
    (python-docx + lxml deepcopy) instead of writing through pandoc.

    Raises on any problem (missing python-docx, unreadable template, no
    usable exemplar paragraph at all) — the caller falls back to the
    pandoc(+cast) path; this function must never be the document's only
    chance to exist.
    """
    from docx import Document  # lazy import: startup-cost
    from docx.enum.text import WD_TAB_ALIGNMENT  # lazy import: startup-cost
    from docx.oxml.ns import qn  # lazy import: startup-cost

    exemplar_source = Document(str(template_docx))
    roles = _learn_template_roles(exemplar_source.paragraphs, qn)
    if not _has_any_exemplar(roles):
        raise DocumentExportError("template has no usable exemplar paragraphs")

    blocks = _parse_markdown_blocks(markdown_body)

    doc = Document(str(template_docx))
    body_el = doc.element.body
    sect_pr = body_el.find(qn("w:sectPr"))
    for child in list(body_el):
        if child is not sect_pr:
            body_el.remove(child)

    counts: dict[str, int] = {}
    for block in blocks:
        exemplar_paragraph = _resolve_exemplar(block["role"], block.get("preamble_index"), roles)
        if exemplar_paragraph is None:
            continue
        new_p = copy.deepcopy(exemplar_paragraph._p)
        _clear_paragraph_runs(new_p, qn)
        _append_block_runs(new_p, block, exemplar_paragraph, qn, WD_TAB_ALIGNMENT)
        if sect_pr is not None:
            sect_pr.addprevious(new_p)
        else:
            body_el.append(new_p)
        counts[block["role"]] = counts.get(block["role"], 0) + 1

    doc.save(str(out))
    logger.info(f"[DocExport] Built from template: paragraphs per role={counts}")


def _docx_with_template(body: str, out: Path, template_docx: Path) -> None:
    """Try the direct template-paragraph builder first; ANY failure falls
    back to the pandoc(+cast) reference-doc path so a build-from-template bug
    can never lose the document."""
    try:
        _build_from_template(body, template_docx, out)
        return
    except Exception as e:  # degrades: falls back to pandoc's reference-doc + cast styling
        logger.warning(f"[DocExport] Template build failed, falling back to pandoc cast: {e}")
    _pandoc_docx(body, out, reference_doc=template_docx)


def _pandoc_docx(body: str, out: Path, *, reference_doc: Path | None = None) -> None:
    """pandoc→docx, then (2026-09-20) cast the user's OWN template's look onto
    the result — see _cast_template_styles. Shared by the plain docx export
    ladder and _soffice_pdf's intermediate docx so a PDF export inherits the
    same cast. Only fires when reference_doc is the caller-supplied file
    (never the built-in compact reference `_pandoc` resolves internally when
    reference_doc is None) — a template the user didn't supply has no look to
    cast. Wrapped defensively here too: the document pandoc already wrote
    must survive even if the cast step is broken or replaced (e.g. by a test)."""
    _pandoc(body, out, "docx", reference_doc=reference_doc)
    if reference_doc is not None:
        try:
            _cast_template_styles(out, reference_doc)
        except Exception as e:  # degrades: pandoc's plain reference-doc styling only
            logger.warning(f"[DocExport] Template style cast failed, leaving pandoc output as-is: {e}")


def _pandoc(body: str, out: Path, fmt: str, *, reference_doc: Path | None = None) -> None:
    """reference_doc: a caller-supplied styling reference (e.g. the user's own
    layout-template attachment) that wins over the built-in compact template
    for docx, or over pandoc's stock odt styling. Caller is responsible for
    only passing one whose suffix matches `fmt` (see export_document)."""
    exe = shutil.which("pandoc")
    if not exe:
        raise DocumentExportError("pandoc is not installed")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "in.md"
        src.write_text(body, encoding="utf-8")
        extra: list[str] = []
        if fmt == "docx":
            ref = reference_doc
            if ref is None:
                try:
                    ref = _reference_docx(tmp)
                except Exception as e:  # degrades: stock pandoc styling (larger, looser)
                    logger.warning(f"[DocExport] Reference template failed, using stock styles: {e}")
                    ref = None
            if ref is not None:
                extra = ["--reference-doc", str(ref)]
        elif fmt == "odt" and reference_doc is not None:
            extra = ["--reference-doc", str(reference_doc)]
        # hard_line_breaks: a composed resume stacks "School — degree" / "Coursework: …"
        # on consecutive lines; plain gfm folds those into one paragraph.
        _run([exe, "-s", "-f", "gfm+hard_line_breaks", "-t", fmt, *extra, "-o", str(out), str(src)])


def _python_docx(body: str, out: Path) -> None:
    """Minimal pure-python renderer: headings, bullets, numbered items, paragraphs."""
    try:
        from docx import Document  # lazy import: startup-cost
    except ImportError as e:
        raise DocumentExportError("python-docx is not installed") from e
    doc = Document()
    for raw in body.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        plain = re.sub(r"(\*\*|__|\*|`)", "", line.strip())
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            doc.add_heading(re.sub(r"(\*\*|__|\*|`)", "", heading.group(2)), level=min(len(heading.group(1)), 4))
        elif re.match(r"^\s*[-*+]\s+", line):
            doc.add_paragraph(re.sub(r"^[-*+]\s+", "", plain), style="List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", line):
            doc.add_paragraph(re.sub(r"^\d+[.)]\s+", "", plain), style="List Number")
        else:
            doc.add_paragraph(plain)
    doc.save(str(out))


def _to_html(body: str) -> str:
    import mistune  # lazy import: startup-cost
    return mistune.create_markdown(hard_wrap=True, plugins=["table", "strikethrough"])(body)


def _convert_docx_to_pdf(docx_path: Path, out: Path, tmp: str, exe: str) -> None:
    """Shared LibreOffice docx->pdf step: an already-built docx (either
    pandoc's output or _build_from_template's) converted headless. Private
    profile dir: a desktop LibreOffice already open on the user's own
    document would otherwise swallow the headless request."""
    _run([exe, f"-env:UserInstallation=file://{tmp}/lo_profile", "--headless",
          "--convert-to", "pdf", "--outdir", tmp, str(docx_path)])
    produced = Path(tmp) / f"{docx_path.stem}.pdf"
    if not produced.exists():
        raise DocumentExportError("LibreOffice produced no PDF")
    shutil.copyfile(produced, out)


def _soffice_pdf(body: str, out: Path, *, reference_doc: Path | None = None) -> None:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise DocumentExportError("LibreOffice is not installed")
    with tempfile.TemporaryDirectory() as tmp:
        docx = Path(tmp) / f"{out.stem}.docx"
        _pandoc_docx(body, docx, reference_doc=reference_doc)
        _convert_docx_to_pdf(docx, out, tmp, exe)


def _pdf_with_template(body: str, out: Path, template_docx: Path) -> None:
    """PDF counterpart of `_docx_with_template`: build the docx directly from
    the template first (falling back to pandoc+cast on ANY failure), then
    convert that single docx to PDF via the existing LibreOffice step."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise DocumentExportError("LibreOffice is not installed")
    with tempfile.TemporaryDirectory() as tmp:
        docx_path = Path(tmp) / f"{out.stem}.docx"
        try:
            _build_from_template(body, template_docx, docx_path)
        except Exception as e:  # degrades: falls back to pandoc's reference-doc + cast styling
            logger.warning(f"[DocExport] Template build failed, falling back to pandoc cast: {e}")
            _pandoc_docx(body, docx_path, reference_doc=template_docx)
        _convert_docx_to_pdf(docx_path, out, tmp, exe)


def _fpdf_pdf(body: str, out: Path) -> None:
    try:
        from fpdf import FPDF  # lazy import: startup-cost
    except ImportError as e:
        raise DocumentExportError("fpdf2 is not installed") from e
    pdf = FPDF()
    pdf.add_page()
    html = _to_html(body)
    font_dir = next((Path(d) for d in _FONT_DIRS if (Path(d) / "DejaVuSans.ttf").exists()), None)
    if font_dir is not None:
        pdf.add_font("DejaVu", "", str(font_dir / "DejaVuSans.ttf"))
        bold = font_dir / "DejaVuSans-Bold.ttf"
        pdf.add_font("DejaVu", "B", str(bold if bold.exists() else font_dir / "DejaVuSans.ttf"))
        oblique = font_dir / "DejaVuSans-Oblique.ttf"
        pdf.add_font("DejaVu", "I", str(oblique if oblique.exists() else font_dir / "DejaVuSans.ttf"))
        pdf.set_font("DejaVu", size=11)
        pdf.write_html(html, font_family="DejaVu")
    else:
        # Core fonts are latin-1 only: degrade the glyphs, keep the document.
        pdf.set_font("Helvetica", size=11)
        pdf.write_html(html.encode("latin-1", "replace").decode("latin-1"))
    pdf.output(str(out))


def _to_txt(body: str) -> str:
    text = re.sub(r"^#{1,6}\s+", "", body, flags=re.M)
    text = re.sub(r"(\*\*|__|`)", "", text)
    return re.sub(r"\[([^\]]+)\]\((?:[^)]+)\)", r"\1", text)


def export_document(md_path: str | Path, fmt: str, *, reference_doc: str | Path | None = None) -> Path:
    """Convert a generated markdown document to `fmt` beside it. Blocking.

    reference_doc (2026-09-20, template-cast documents): a caller-supplied
    styling reference — e.g. a second attachment the user supplied as a
    layout template — used INSTEAD of the built-in compact reference when its
    format matches the target (.docx for docx/pdf, .odt for odt). A template
    in another format (pdf, md, txt, …) contributes structure text only at
    the compose_from_material stage — it is not an export error here, just
    ignored for styling. A malformed/corrupt user template that makes pandoc
    fail is retried once with the built-in reference (see the docx/odt/pdf
    ladders below) so it can never lose the document.
    """
    fmt = (fmt or "").lower().lstrip(".")
    if fmt not in EXPORT_FORMATS:
        raise DocumentExportError(f"unsupported format: {fmt!r}")
    src = Path(md_path)
    body = strip_frontmatter(src.read_text(encoding="utf-8"))
    out = _versioned(src.with_suffix(f".{fmt}"))

    if fmt == "html":
        out.write_text(
            f"<!doctype html><meta charset=\"utf-8\"><title>{src.stem}</title>\n{_to_html(body)}",
            encoding="utf-8",
        )
        return out
    if fmt == "txt":
        out.write_text(_to_txt(body), encoding="utf-8")
        return out

    ref_path = Path(reference_doc) if reference_doc else None
    if ref_path is not None and not ref_path.exists():
        logger.warning(f"[DocExport] Reference template not found, ignoring: {ref_path}")
        ref_path = None

    def _matching_ref(target_suffix: str) -> Path | None:
        return ref_path if (ref_path is not None and ref_path.suffix.lower() == target_suffix) else None

    # Ladders below only prepend a "use the user's template" step when one
    # actually applies (matching suffix) — otherwise they are byte-identical
    # to the pre-template ladder, so the no-template case (still the common
    # one) never pays for a duplicate identical attempt.
    _docx_ref = _matching_ref(".docx")
    _odt_ref = _matching_ref(".odt")

    docx_steps = []
    if _docx_ref is not None:
        # Direct template build first (2026-09-20) — falls back to
        # pandoc(+cast) internally on ANY failure, see _docx_with_template.
        docx_steps.append(lambda: _docx_with_template(body, out, _docx_ref))
    docx_steps += [lambda: _pandoc(body, out, "docx"), lambda: _python_docx(body, out)]

    odt_steps = []
    if _odt_ref is not None:
        odt_steps.append(lambda: _pandoc(body, out, "odt", reference_doc=_odt_ref))
    odt_steps.append(lambda: _pandoc(body, out, "odt"))

    pdf_steps = []
    if _docx_ref is not None:
        pdf_steps.append(lambda: _pdf_with_template(body, out, _docx_ref))
    pdf_steps += [lambda: _soffice_pdf(body, out), lambda: _fpdf_pdf(body, out)]

    ladder = {
        "docx": docx_steps,
        "odt": odt_steps,
        "rtf": [lambda: _pandoc(body, out, "rtf")],
        "pdf": pdf_steps,
    }[fmt]
    errors: list[str] = []
    for step in ladder:
        try:
            step()
            if out.exists() and out.stat().st_size > 0:
                return out
            errors.append("converter produced an empty file")
        except (DocumentExportError, subprocess.TimeoutExpired, OSError) as e:
            errors.append(str(e))
            logger.warning(f"[DocExport] {fmt} converter failed, trying next: {e}")
    raise DocumentExportError(f"could not write .{fmt}: " + "; ".join(errors))
